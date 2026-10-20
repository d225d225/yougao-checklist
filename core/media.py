"""上傳檔案處理：圖片壓縮、HEIC 轉檔、轉成 Claude 可讀的內容區塊。"""
from __future__ import annotations

import base64
import io
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image, ImageOps

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
except ImportError:  # 沒裝 pillow-heif 時，HEIC 會被當成「無法讀取」
    pass

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp", ".gif", ".bmp"}
MAX_PDF_BYTES = 25 * 1024 * 1024


@dataclass
class Upload:
    name: str
    data: bytes
    mime: str
    kind: str = field(init=False)  # image / pdf / docx / other

    def __post_init__(self):
        ext = Path(self.name).suffix.lower()
        if ext in IMAGE_EXT or self.mime.startswith("image/"):
            self.kind = "image"
        elif ext == ".pdf" or self.mime == "application/pdf":
            self.kind = "pdf"
        elif ext == ".docx":
            self.kind = "docx"
        else:
            self.kind = "other"


def open_image(data: bytes) -> Image.Image:
    img = Image.open(io.BytesIO(data))
    img = ImageOps.exif_transpose(img)  # 手機照片轉正
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    return img


def to_jpeg(data: bytes, max_side: int, quality: int) -> bytes:
    img = open_image(data)
    img.thumbnail((max_side, max_side))
    out = io.BytesIO()
    img.save(out, format="JPEG", quality=quality, optimize=True)
    return out.getvalue()


def pdf_first_page_jpeg(data: bytes, max_side: int = 1200, quality: int = 80) -> bytes | None:
    """把 PDF 第一頁轉成 JPEG（給獎狀、公文當照片用）。失敗時回傳 None。"""
    try:
        import pypdfium2 as pdfium

        page = pdfium.PdfDocument(data)[0]
        img = page.render(scale=2).to_pil().convert("RGB")
        img.thumbnail((max_side, max_side))
        out = io.BytesIO()
        img.save(out, format="JPEG", quality=quality, optimize=True)
        return out.getvalue()
    except Exception:
        return None


def photo_jpeg(up: "Upload", max_side: int = 1200, quality: int = 80) -> bytes | None:
    """可放進相片欄的 JPEG：圖片直接壓縮，PDF 取第一頁；其他格式回傳 None。"""
    if up.kind == "image":
        return to_jpeg(up.data, max_side, quality)
    if up.kind == "pdf":
        return pdf_first_page_jpeg(up.data, max_side, quality)
    return None


def docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append("｜".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def content_blocks(uploads: list[Upload]) -> tuple[list[dict], list[str]]:
    """把上傳檔案轉成 Claude 內容區塊。回傳 (blocks, warnings)。

    每個檔案前面加一行「檔案 #i：檔名」，讓 AI 回報時能用編號指認檔案。
    """
    blocks: list[dict] = []
    warnings: list[str] = []
    for i, up in enumerate(uploads):
        label = {"type": "text", "text": f"檔案 #{i}：{up.name}"}
        try:
            if up.kind == "image":
                jpg = to_jpeg(up.data, 1280, 80)
                blocks += [label, {"type": "image", "source": {
                    "type": "base64", "media_type": "image/jpeg",
                    "data": base64.standard_b64encode(jpg).decode("ascii")}}]
            elif up.kind == "pdf":
                if len(up.data) > MAX_PDF_BYTES:
                    warnings.append(f"{up.name} 超過 25MB，AI 不讀取此檔")
                    continue
                blocks += [label, {"type": "document", "source": {
                    "type": "base64", "media_type": "application/pdf",
                    "data": base64.standard_b64encode(up.data).decode("ascii")}}]
            elif up.kind == "docx":
                blocks += [label, {"type": "text", "text": docx_text(up.data)[:20000]}]
            else:
                warnings.append(f"{up.name} 格式不支援，AI 不讀取此檔（仍會存到雲端）")
        except Exception as e:  # 單一檔案壞掉不影響其他檔案
            warnings.append(f"{up.name} 讀取失敗：{e}")
    return blocks, warnings


def safe_name(s: str, limit: int = 40) -> str:
    bad = '\\/:*?"<>|\n\r\t'
    s = "".join("_" if ch in bad else ch for ch in s).strip()
    return s[:limit] or "未命名"
